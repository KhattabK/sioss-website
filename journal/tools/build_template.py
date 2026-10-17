"""Builds journal/files/SJSS-paper-template.docx: the SIOSS-branded manuscript template for SJSS.

Run with python-docx installed:  python build_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SITE = HERE.parent.parent
OUT = HERE.parent / "files" / "SJSS-paper-template.docx"
LOGO = SITE / "pt" / "logo-on-light.png"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1A, 0x6F, 0xC4)
GOLD = RGBColor(0xC9, 0xA2, 0x4A)
GREY = RGBColor(0x5B, 0x66, 0x78)
HEAD, BODY = "Calibri", "Cambria"


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def para_border(p, edge="bottom", sz=8, color="C9A24A", space=4):
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    pb.append(el); pPr.append(pb)


def set_font(style_or_run, name, size=None, bold=None, italic=None, color=None):
    f = style_or_run.font
    f.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color


def style(doc, name, base="Normal", font=BODY, size=10, bold=False, italic=False, color=None,
          before=0, after=4, align=None, line=1.15, keep=False):
    st = doc.styles[name] if name in [s.name for s in doc.styles] else doc.styles.add_style(name, 1)
    if name != "Normal":
        st.base_style = doc.styles[base]
    set_font(st, font, size, bold, italic, color)
    pf = st.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if align is not None: pf.alignment = align
    pf.keep_with_next = keep
    return st


def columns(section, n, space_cm=0.8):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sectPr.append(cols)
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(int(space_cm * 567)))


def page_field(run):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


doc = Document()
# --- styles
style(doc, "Normal", font=BODY, size=10, after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
style(doc, "SJSS Title", font=HEAD, size=20, bold=True, color=NAVY, before=6, after=8, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT)
style(doc, "SJSS Authors", font=HEAD, size=11.5, bold=True, color=BLUE, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
style(doc, "SJSS Affiliation", font=HEAD, size=8.5, color=GREY, after=1, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)
style(doc, "SJSS Abstract", font=BODY, size=9.5, after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
style(doc, "SJSS Small", font=HEAD, size=8, color=GREY, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
h1 = style(doc, "Heading 1", font=HEAD, size=12, bold=True, color=NAVY, before=10, after=4, keep=True, align=WD_ALIGN_PARAGRAPH.LEFT)
h2 = style(doc, "Heading 2", font=HEAD, size=10.5, bold=True, color=BLUE, before=6, after=3, keep=True, align=WD_ALIGN_PARAGRAPH.LEFT)
for h in (h1, h2):
    h.font.italic = False
style(doc, "Caption", font=HEAD, size=8.5, color=NAVY, before=4, after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
style(doc, "SJSS Reference", font=BODY, size=8.5, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
doc.styles["SJSS Reference"].paragraph_format.left_indent = Cm(0.6)
doc.styles["SJSS Reference"].paragraph_format.first_line_indent = Cm(-0.6)

# --- page setup
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(0.9), Cm(0.9)

# header: journal line + logo
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = hp.add_run("SJSS"); set_font(r, HEAD, 9, True, color=NAVY)
r = hp.add_run("  ·  The SIOSS Journal of Sport Science  ·  ISSN 2519-6634 (Online)"); set_font(r, HEAD, 8, color=GREY)
r = hp.add_run("\t"); hp.add_run().add_picture(str(LOGO), height=Cm(0.75))
pPr = hp._p.get_or_add_pPr()
tabs = OxmlElement("w:tabs")
for pos, kind in ((4680, "clear"), (9360, "clear"), (int(17.4 * 567), "right")):  # drop the Header style's default stops
    t = OxmlElement("w:tab"); t.set(qn("w:val"), kind); t.set(qn("w:pos"), str(pos)); tabs.append(t)
pPr.append(tabs)
para_border(hp, sz=6, color="C9A24A", space=4)

# footer: open access + page number
fp = sec.footer.paragraphs[0]
r = fp.add_run("Open access · Published by the Swedish International Organization of Sport Science (SIOSS), Linköping, Sweden · www.sioss.se/journal")
set_font(r, HEAD, 7.5, color=GREY)
fp.add_run("\t")
r = fp.add_run(); set_font(r, HEAD, 9, True, color=NAVY); page_field(r)
fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), 2)  # right
para_border(fp, edge="top", sz=4, color="D0D7E2", space=4)

# --- first-page banner
t = doc.add_table(rows=1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
left, right = t.rows[0].cells
left.width, right.width = Cm(11.4), Cm(6.0)
for c in (left, right):
    shade(c, "0B1F3A"); cell_borders(c)
p = left.paragraphs[0]
r = p.add_run("ORIGINAL RESEARCH"); set_font(r, HEAD, 8, True, color=RGBColor(0xF5, 0xB8, 0x00))
r.font.spacing = Pt(1.5) if hasattr(r.font, "spacing") else None
p.paragraph_format.space_before = Pt(4)
p = left.add_paragraph(); r = p.add_run("The SIOSS Journal of Sport Science"); set_font(r, HEAD, 13, True, color=RGBColor(0xFF, 0xFF, 0xFF))
p.paragraph_format.space_after = Pt(4)
p = right.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("Vol. X, Issue X, 20XX · pp. X–X"); set_font(r, HEAD, 8.5, color=RGBColor(0xCF, 0xDC, 0xEC))
p.paragraph_format.space_before = Pt(6)
p = right.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("[Filled in by the editorial office]"); set_font(r, HEAD, 7.5, italic=True, color=RGBColor(0x9F, 0xB0, 0xC6))
p.paragraph_format.space_after = Pt(4)

doc.add_paragraph("Paper Title: Clear, Specific and Informative (Maximum 20 Words)", style="SJSS Title")

p = doc.add_paragraph(style="SJSS Authors")
for i, name in enumerate(("First Author", "Second Author", "Third Author"), 1):
    p.add_run(name)
    sup = p.add_run(str(i) + ("*" if i == 1 else "")); sup.font.superscript = True
    if i < 3: p.add_run(", ")
for i, aff in enumerate(("Department, University, City, Country", "Department, University, City, Country",
                         "Department, Institution, City, Country"), 1):
    p = doc.add_paragraph(style="SJSS Affiliation")
    sup = p.add_run(str(i)); sup.font.superscript = True
    p.add_run(" " + aff)
p = doc.add_paragraph(style="SJSS Affiliation")
r = p.add_run("*Corresponding author: "); r.bold = True
p.add_run("name@university.edu  ·  ORCID: 0000-0000-0000-0000")
p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph(style="SJSS Small")
p.add_run("Received: DD Month 20XX   ·   Accepted: DD Month 20XX   ·   Published: DD Month 20XX   ").italic = True
r = p.add_run("(filled in by the editorial office)"); r.italic = True; r.font.color.rgb = GOLD
p.paragraph_format.space_after = Pt(8)

# abstract box
box = doc.add_table(rows=1, cols=1); box.alignment = WD_TABLE_ALIGNMENT.CENTER
c = box.rows[0].cells[0]; c.width = Cm(17.4)
shade(c, "EEF3F9"); cell_borders(c, left=(24, "1A6FC4"))
p = c.paragraphs[0]; r = p.add_run("Abstract"); set_font(r, HEAD, 10.5, True, color=NAVY)
p.paragraph_format.space_before = Pt(4)
p = c.add_paragraph(style="SJSS Abstract")
for label, text in (("Background: ", "State the problem and the aim of the study in one or two sentences. "),
                    ("Methods: ", "Describe the participants, design, measurements and statistical analysis. "),
                    ("Results: ", "Give the main findings with key numbers (e.g. mean ± SD, effect size, p-value). "),
                    ("Conclusion: ", "State the main conclusion and its practical meaning for sport, physical education or health. "
                                     "The abstract should not exceed 250 words and should contain no references.")):
    r = p.add_run(label); r.bold = True; r.font.color.rgb = NAVY
    p.add_run(text)
p = c.add_paragraph(style="SJSS Abstract")
r = p.add_run("Keywords: "); r.bold = True; r.font.color.rgb = NAVY
p.add_run("keyword one; keyword two; keyword three; keyword four; keyword five (3–6 keywords)")
p.paragraph_format.space_after = Pt(6)

# --- two-column body
body = doc.add_section(WD_SECTION.CONTINUOUS)
columns(body, 2)


def h(text, lvl=1):
    doc.add_paragraph(text, style=f"Heading {lvl}")


def para(text):
    doc.add_paragraph(text)


h("1. Introduction")
para("Use this template by replacing the placeholder text with your own; the styles for the title, authors, headings, "
     "body text, captions and references are already set. Do not change margins, columns or fonts. Write in clear, "
     "concise English and number the main sections as shown (1., 2., 3. …) and the subsections as 2.1, 2.2 and so on.")
para("The introduction presents the background of the problem, summarises the most relevant previous research "
     "(Author, 2020; Author & Author, 2022), identifies the gap that the study addresses and ends with the aim, the "
     "research questions or the hypotheses.")
h("2. Methods")
h("2.1 Participants", 2)
para("Describe the sample: number, sex, age (mean ± SD), training level and inclusion and exclusion criteria. State that "
     "all participants (or their guardians) gave written informed consent and name the ethics committee that approved "
     "the study, with the approval number.")
h("2.2 Procedures and measurements", 2)
para("Describe the study design, the tests and instruments, the intervention (duration, frequency, intensity) and "
     "the testing conditions in enough detail for another researcher to repeat the study.")
h("2.3 Statistical analysis", 2)
para("Name the statistical software and version, the tests used, the significance level (e.g. p < 0.05) and how effect "
     "sizes were calculated.")
h("3. Results")
para("Present the results in a logical order using text, tables and figures without repeating the same data. "
     "Refer to every table and figure in the text (Table 1, Figure 1).")
cap = doc.add_paragraph(style="Caption")
r = cap.add_run("Table 1. "); r.bold = True
cap.add_run("Characteristics of the participants (mean ± SD).")
tb = doc.add_table(rows=4, cols=3); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
rows = [("Variable", "Group A (n = 20)", "Group B (n = 20)"), ("Age (years)", "15.2 ± 1.1", "15.4 ± 0.9"),
        ("Height (cm)", "168.4 ± 6.2", "167.9 ± 5.8"), ("VO₂max (ml/kg/min)", "48.1 ± 4.3", "47.6 ± 4.9")]
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cl = tb.rows[i].cells[j]
        cl.text = ""
        pr = cl.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        rr = pr.add_run(val); set_font(rr, HEAD, 8, bold=(i == 0), color=(RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None))
        if i == 0: shade(cl, "0B1F3A")
        elif i % 2 == 0: shade(cl, "F2F5F9")
        cell_borders(cl, bottom=(4, "D0D7E2"))
doc.add_paragraph().paragraph_format.space_after = Pt(2)
fig = doc.add_table(rows=1, cols=1); fig.alignment = WD_TABLE_ALIGNMENT.CENTER
fc = fig.rows[0].cells[0]; shade(fc, "F7F9FC"); cell_borders(fc, top=(6, "B8C4D6"), bottom=(6, "B8C4D6"), left=(6, "B8C4D6"), right=(6, "B8C4D6"))
fp2 = fc.paragraphs[0]; fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3): fp2.add_run().add_break()
rr = fp2.add_run("[Insert figure here: high resolution, at least 300 dpi]"); set_font(rr, HEAD, 8, italic=True, color=GREY)
for _ in range(3): fp2.add_run().add_break()
cap = doc.add_paragraph(style="Caption")
r = cap.add_run("Figure 1. "); r.bold = True
cap.add_run("Short, self-explanatory caption. Explain all abbreviations and symbols used in the figure.")
h("4. Discussion")
para("Interpret the main findings, compare them with previous studies, explain possible mechanisms, and discuss the "
     "practical applications for coaches, teachers or clinicians. State the limitations of the study honestly and "
     "suggest directions for future research.")
h("5. Conclusion")
para("Summarise the answer to the research question in a few sentences, without repeating the results in detail.")
h("Acknowledgements")
para("Thank the people and institutions who helped with the study (optional).")
h("Conflict of interest and funding")
para("Declare any conflict of interest, or write: “The authors declare no conflict of interest.” Name any funding "
     "body and grant number.")
h("References")
p = doc.add_paragraph(style="SJSS Small")
r = p.add_run("Use the APA style (7th edition). List references alphabetically; in the text cite as (Author, Year). "
              "Include the DOI wherever available.")
r.italic = True
p.paragraph_format.space_after = Pt(4)
for ref in (
    "Author, A. A., & Author, B. B. (2022). Title of the journal article in sentence case. Journal Name in Title Case, 12(3), 145–160. https://doi.org/10.xxxx/xxxxx",
    "Author, C. C. (2020). Title of the book in sentence case (2nd ed.). Publisher.",
    "Author, D. D., Author, E. E., & Author, F. F. (2021). Title of the chapter. In G. Editor (Ed.), Title of the book (pp. 45–67). Publisher.",
    "Organisation Name. (2023). Title of the web page or report. https://www.example.org/page",
):
    doc.add_paragraph(ref, style="SJSS Reference")

# --- end-of-template checklist (one column)
end = doc.add_section(WD_SECTION.NEW_PAGE)
columns(end, 1)
p = doc.add_paragraph(style="SJSS Title"); p.add_run("Before you submit (delete this page)").font.size = Pt(14)
para_border(p, sz=8)
for item in (
    "The paper is written in English and does not exceed 10 pages in this template.",
    "The title, all authors, affiliations and the corresponding author’s email are complete.",
    "The abstract has no more than 250 words, followed by 3–6 keywords.",
    "Every table and figure is numbered, has a caption and is cited in the text.",
    "Every in-text citation appears in the reference list (APA 7th edition) and vice versa.",
    "Ethical approval and informed consent are stated for studies with human participants.",
    "The paper has not been published or submitted for publication elsewhere.",
    "Save the final file as PDF, pay the publication fee at www.sioss.se/journal/instructions-ar.html and upload the PDF on the page that opens after payment.",
):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rr = q.add_run("☐  "); set_font(rr, "Segoe UI Symbol", 11, color=BLUE)
    q.add_run(item)
q = doc.add_paragraph(style="SJSS Small"); q.paragraph_format.space_before = Pt(10)
q.add_run("Questions: hello@sioss.se  ·  WhatsApp +46 70 718 11 37")

core = doc.core_properties
core.title, core.author, core.subject = "SJSS paper template", "SIOSS – The SIOSS Journal of Sport Science", "Manuscript template"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT)
